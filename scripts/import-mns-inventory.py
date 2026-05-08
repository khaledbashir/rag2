#!/usr/bin/env python3
"""
Universal importer for the multi-sponsor team sheets in 2026 ANC MLB Inventory List.

For each team sheet, parses the per-position sponsor occupancy columns and creates
one mediaPlacement record per (game, sponsor, position) combination.

Idempotent: skips any (gameDate, homeTeamId, sponsorId, placementType) tuple already
in the CRM.

Run:  python3 /root/rag2/scripts/import-mns-inventory.py
"""
import json
import os
import re
import sys
import time
import urllib.request

import openpyxl

TWENTY_BASE = "https://abc-twenty.izcgmb.easypanel.host"
TWENTY_API_KEY = os.environ.get("TWENTY_API_KEY") or (
    "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9."
    "eyJzdWIiOiJkM2ZiYzI5YS1hNjM1LTQ4YjctOWQ2ZS0yNTA5NDE2NzdmZDAiLCJ0eXBlIjoiQVBJX0tFWSIsIndvcmtzcGFjZUlkIjoiZDNmYmMyOWEtYTYzNS00OGI3LTlkNmUtMjUwOTQxNjc3ZmQwIiwiaWF0IjoxNzc0ODEwNDkyLCJleHAiOjQ5Mjg0MTA0ODcsImp0aSI6IjYxMGEzMWEzLTJhMDgtNDM5MC1iMTU1LTFkN2M3NzY5Y2QxOSJ9."
    "nzknS-bBNuf7y3LUCv2xEa5-9xuJNHBK3GalJwWK3eA"
)

# Reuse team map from kumho importer
sys.path.insert(0, os.path.dirname(__file__))
from importlib import import_module
TEAM_MAP = import_module("import-mns-kumho").TEAM_MAP

# Sponsor name (any case, with/without trailing words like "Tire", "Tea") -> Company UUID
SPONSOR_MAP = {
    "kumho":                   "6c6b2965-b870-40a4-a9b8-568eee485257",  # Kumho Tire
    "kumho tire":              "6c6b2965-b870-40a4-a9b8-568eee485257",
    "tire rack":               "9a48b042-33cd-47ff-a38b-7c789a1fb3ca",  # Tire Rack.com
    "tirerack":                "9a48b042-33cd-47ff-a38b-7c789a1fb3ca",
    "bigelow":                 "4f8b25ce-57e7-4fa4-909b-e953a78b4172",  # Bigelow Tea
    "bigelow tea":             "4f8b25ce-57e7-4fa4-909b-e953a78b4172",
    "hankook":                 "a2522508-e708-4d91-8ee2-75e200592506",  # Hankook Tire America Corp.
    "hankook tire":            "a2522508-e708-4d91-8ee2-75e200592506",
    "discount tire":           "7551ac58-4e04-4343-bee0-f5686a2da02c",  # Discount Tire
    "lucas oil":               "5f120484-89a6-4a29-8de5-8fe5709c1f5d",  # Lucas Oil Products
    "lumber liquidators":      "d9f466db-2ed3-424b-8180-86fe9b126faa",
    "lumber/cabinets":         "d9f466db-2ed3-424b-8180-86fe9b126faa",
    "lumber":                  "d9f466db-2ed3-424b-8180-86fe9b126faa",
    "kenda":                   "8ae8221c-e0b2-44d1-9f51-0a5e7c5cc428",
    "kenda tire":              "8ae8221c-e0b2-44d1-9f51-0a5e7c5cc428",
    "ctg":                     "34a8eeee-c847-4f9c-af78-412c46448fc8",  # Cabinets To Go
    "cabinets to go":          "34a8eeee-c847-4f9c-af78-412c46448fc8",
    "cabinets":                "34a8eeee-c847-4f9c-af78-412c46448fc8",
    "lg":                      "14d1e93d-bd7e-42a3-a75b-5cc5c34c4a51",  # LG Electronics
    "lg electronics":          "14d1e93d-bd7e-42a3-a75b-5cc5c34c4a51",
    "nhtsa":                   "2beca1a4-71c1-4390-a7d6-f8d56d7ffc07",
    "authority":               "67ee1571-53da-4911-aab9-9ec5eb833dc8",
    "authority brands":        "67ee1571-53da-4911-aab9-9ec5eb833dc8",
    "l and s":                 None,  # ambiguous shorthand — skip
    "l & s":                   None,
}

# Sheet name -> home team id (overrides per-row Home Team col, since per-team sheets
# have just-the-team rows and the Home Team col may be inconsistent)
SHEET_TO_HOME = {
    "Arizona Diamondbacks":   TEAM_MAP["Arizona Diamondbacks"],
    "Atlanta Braves":         TEAM_MAP["Atlanta Braves"],
    "Boston Red Sox":         TEAM_MAP["Boston Red Sox"],
    "Cincinnati Reds":        TEAM_MAP["Cincinnati Reds"],
    "Colorado Rockies":       TEAM_MAP["Colorado Rockies"],
    "Chicago White Sox":      TEAM_MAP["Chicago White Sox"],
    "Detroit Tigers":         TEAM_MAP["Detroit Tigers"],
    "Kansas City Royals":     TEAM_MAP["Kansas City Royals"],
    "Los Angeles Angels":     TEAM_MAP["Los Angeles Angels"],
    "Miami Marlins":          TEAM_MAP["Miami Marlins"],
    "Milwaukee Brewers":      TEAM_MAP["Milwaukee Brewers"],
    "NY Mets ST":             TEAM_MAP["New York Mets"],
    "NY Yankees":             TEAM_MAP["New York Yankees"],
    "Philadelphia Phillies ST": TEAM_MAP["Philadelphia Phillies"],
    "NY Mets":                TEAM_MAP["New York Mets"],
    "Philadelphia Phillies":  TEAM_MAP["Philadelphia Phillies"],
    "Seattle Mariners":       TEAM_MAP["Seattle Mariners"],
    "St. Louis Cardinals":    TEAM_MAP["St. Louis Cardinals"],
    "Tampa Bay Rays":         TEAM_MAP["Tampa Bay Rays"],
    "Toronto Blue Jays":      TEAM_MAP["Toronto Blue Jays"],
    "WAS Nationals":          TEAM_MAP["Washington Nationals"],
}


def gql(query, variables=None):
    body = json.dumps({"query": query, "variables": variables or {}}).encode("utf-8")
    req = urllib.request.Request(
        f"{TWENTY_BASE}/graphql", data=body,
        headers={"Authorization": f"Bearer {TWENTY_API_KEY}", "Content-Type": "application/json"},
        method="POST")
    with urllib.request.urlopen(req, timeout=60) as resp:
        d = json.loads(resp.read().decode("utf-8"))
    if d.get("errors"):
        raise RuntimeError(f"GraphQL: {json.dumps(d['errors'])[:500]}")
    return d["data"]


# ---- header parsing ----

def normalize_header(s):
    return re.sub(r"\s+", " ", str(s or "").replace("\n", " ").strip()).lower()


def find_header_row(ws):
    for r in range(1, min(ws.max_row + 1, 12)):
        row = [normalize_header(ws.cell(r, c).value) for c in range(1, min(ws.max_column + 1, 30))]
        if "game date" in row and "home team" in row:
            return r
    return None


POSITION_HEADER_PATTERNS = [
    (re.compile(r"homeplate.*1.*full|home plate.*1.*full", re.I), "HOMEPLATE_FULL_1"),
    (re.compile(r"homeplate.*2.*full|home plate.*2.*full", re.I), "HOMEPLATE_FULL_2"),
    (re.compile(r"homeplate.*3.*full|home plate.*3.*full", re.I), "HOMEPLATE_FULL_2"),  # treat 3 as second full alias
    (re.compile(r"homeplate.*1.*half|home plate.*1.*half", re.I), "HOMEPLATE_HALF_1"),
    (re.compile(r"homeplate.*2.*half|home plate.*2.*half", re.I), "HOMEPLATE_HALF_2"),
    (re.compile(r"small home plate.*1", re.I),                   "SMALL_HOMEPLATE_1"),
    (re.compile(r"small home plate.*2", re.I),                   "SMALL_HOMEPLATE_2"),
    (re.compile(r"homeplate.*right.*1|home plate.*right.*1", re.I), "HOMEPLATE_FULL_1"),
    (re.compile(r"homeplate.*right.*2|home plate.*right.*2", re.I), "HOMEPLATE_FULL_2"),
    (re.compile(r"homeplate.*right.*3|home plate.*right.*3", re.I), "HOMEPLATE_HALF_1"),
    (re.compile(r"homeplate #?1", re.I),                          "HOMEPLATE_FULL_1"),
    (re.compile(r"homeplate #?2", re.I),                          "HOMEPLATE_FULL_2"),
    (re.compile(r"home ?plate position 1|homeplate position 1|homeplate 1", re.I), "HOMEPLATE_FULL_1"),
    (re.compile(r"home ?plate position 2|homeplate position 2|homeplate 2", re.I), "HOMEPLATE_FULL_2"),
    (re.compile(r"home ?plate position 3|homeplate position 3|homeplate 3", re.I), "HOMEPLATE_HALF_1"),
    (re.compile(r"home plate \(half\)|homeplate \(half\)", re.I), "HOMEPLATE_HALF_1"),
    (re.compile(r"home plate 4|homeplate 4", re.I),               "HOMEPLATE_HALF_2"),
    (re.compile(r"led dugout lip.*position 1", re.I),             "OTHER"),
    (re.compile(r"led dugout lip.*position 2", re.I),             "OTHER"),
    (re.compile(r"outfield led", re.I),                           "OTHER"),
    (re.compile(r"^homeplate$", re.I),                            "HOMEPLATE_FULL_1"),
    (re.compile(r"^home plate", re.I),                            "HOMEPLATE_FULL_1"),
]


def detect_position_columns(header_cells):
    """Return list of (position_col_idx, inning_col_idx, placementType) tuples.

    A position column has a header that matches a placement pattern AND is followed
    by a column whose header is 'Inning Scheduled' (variants ok).
    Stops scanning once a 'Game Count' or 'KEY' column is found.
    """
    pairs = []
    n = len(header_cells)
    i = 0
    while i < n:
        h = header_cells[i]
        if "game count" in h or h == "key":
            break
        # match position header
        match_type = None
        for pat, ptype in POSITION_HEADER_PATTERNS:
            if pat.search(h):
                match_type = ptype
                break
        if match_type and i + 1 < n and "inning" in header_cells[i + 1]:
            pairs.append((i, i + 1, match_type))
            i += 2
            continue
        i += 1
    return pairs


def find_data_columns(header_cells):
    out = {}
    for c in ("game date", "time (est)", "home team", "away team",
              "home tv network", "away tv network", "national tv network"):
        for i, h in enumerate(header_cells):
            if h == c:
                out[c] = i
                break
    return out


def normalize_team(name):
    if not name:
        return None
    return str(name).strip()


def normalize_sponsor(cell):
    if cell is None:
        return None
    s = str(cell).strip().lower()
    if not s or s in ("national game", "national", "ng"):
        return ("NATIONAL_GAME", None)  # signal national-game row, no sponsor
    # exact map
    if s in SPONSOR_MAP:
        return ("SPONSOR", SPONSOR_MAP[s])
    # fuzzy contains
    for key, sid in SPONSOR_MAP.items():
        if sid and key in s:
            return ("SPONSOR", sid)
    return ("UNKNOWN", s)


def existing_keys():
    """Build a set of (gameDate, sponsorId, homeTeamId, placementType) tuples
    already in CRM, so we don't double-import."""
    out = set()
    after = None
    while True:
        var = {"first": 200}
        if after:
            var["after"] = after
        d = gql("""
            query All($first: Int, $after: String) {
              mediaPlacements(first: $first, after: $after) {
                edges {
                  node { gameDate sponsorId homeTeamId placementType }
                  cursor
                }
                pageInfo { hasNextPage endCursor }
              }
            }
        """, var)
        edges = d["mediaPlacements"]["edges"]
        for e in edges:
            n = e["node"]
            d_str = (n.get("gameDate") or "")[:10]
            out.add((d_str, n.get("sponsorId"), n.get("homeTeamId"), n.get("placementType")))
        if not d["mediaPlacements"]["pageInfo"]["hasNextPage"]:
            break
        after = d["mediaPlacements"]["pageInfo"]["endCursor"]
    return out


def create_placement(row):
    d = gql("""
        mutation Create($data: MediaPlacementCreateInput!) {
          createMediaPlacement(data: $data) { id }
        }
    """, {"data": row})
    return d["createMediaPlacement"]["id"]


def parse_sheet(ws, sheet_name, existing, log):
    home_team_id = SHEET_TO_HOME.get(sheet_name)
    if not home_team_id:
        log.append((sheet_name, "no home-team mapping for sheet"))
        return 0, 0, 0

    header_row = find_header_row(ws)
    if not header_row:
        log.append((sheet_name, "no header row found"))
        return 0, 0, 0

    headers = [normalize_header(ws.cell(header_row, c).value) for c in range(1, ws.max_column + 1)]
    cols = find_data_columns(headers)
    pairs = detect_position_columns(headers)
    if not pairs:
        log.append((sheet_name, f"no position pairs detected in headers: {headers[:30]}"))
        return 0, 0, 0

    created = 0
    skipped = 0
    unknown_sponsors = set()

    for r in range(header_row + 1, ws.max_row + 1):
        date_cell = ws.cell(r, cols.get("game date", -1) + 1).value if "game date" in cols else None
        if not date_cell:
            continue
        try:
            date_str = date_cell.strftime("%Y-%m-%d")
        except Exception:
            continue
        season = date_cell.year if hasattr(date_cell, "year") else 2026

        time_cell = ws.cell(r, cols.get("time (est)", -1) + 1).value if "time (est)" in cols else None
        time_str = None
        if time_cell:
            try:
                time_str = time_cell.strftime("%H:%M:%S")
            except Exception:
                time_str = str(time_cell)

        away_cell = ws.cell(r, cols.get("away team", -1) + 1).value if "away team" in cols else None
        away_team_id = TEAM_MAP.get(normalize_team(away_cell))

        home_net = ws.cell(r, cols.get("home tv network", -1) + 1).value if "home tv network" in cols else None
        away_net = ws.cell(r, cols.get("away tv network", -1) + 1).value if "away tv network" in cols else None
        national_net = ws.cell(r, cols.get("national tv network", -1) + 1).value if "national tv network" in cols else None

        # If a national game (national network present) and per-position cells empty, skip
        for pos_col, inning_col, ptype in pairs:
            sponsor_cell = ws.cell(r, pos_col + 1).value
            inning_cell = ws.cell(r, inning_col + 1).value
            if not sponsor_cell:
                continue
            kind, payload = normalize_sponsor(sponsor_cell)
            if kind == "NATIONAL_GAME":
                continue
            if kind == "UNKNOWN":
                unknown_sponsors.add(payload)
                continue
            sponsor_id = payload
            key = (date_str, sponsor_id, home_team_id, ptype)
            if key in existing:
                skipped += 1
                continue
            try:
                exec_status = "SCHEDULED"
                if national_net and not sponsor_cell:
                    exec_status = "NATIONAL_GAME"
                payload_data = {
                    "gameDate": date_str,
                    "gameTime": time_str,
                    "homeTeamId": home_team_id,
                    "sponsorId": sponsor_id,
                    "homeTVNetwork": str(home_net).strip() if home_net else None,
                    "awayTVNetwork": str(away_net).strip() if away_net else None,
                    "nationalTVNetwork": str(national_net).strip() if national_net else None,
                    "inningScheduled": str(inning_cell).strip() if inning_cell else None,
                    "executionStatus": exec_status,
                    "league": "MLB",
                    "season": season,
                    "placementType": ptype,
                    "nielsenVerified": False,
                    "sponsorGameCount": 1,
                }
                if away_team_id:
                    payload_data["awayTeamId"] = away_team_id
                create_placement(payload_data)
                created += 1
                existing.add(key)
                if created % 25 == 0:
                    print(f"    {sheet_name}: {created} created so far")
                time.sleep(0.04)
            except Exception as e:
                log.append((sheet_name, f"row {r} pos {ptype}: {str(e)[:160]}"))

    if unknown_sponsors:
        log.append((sheet_name, f"unmapped sponsors: {sorted(unknown_sponsors)}"))
    return created, skipped, len(unknown_sponsors)


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else "/root/rag2/2026 ANC MLB Inventory List_5_8.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True)
    print(f"Loading existing mediaPlacements for idempotency...")
    existing = existing_keys()
    print(f"  {len(existing)} existing")

    log = []
    total_created = 0
    total_skipped = 0
    for sname in wb.sheetnames:
        if sname in ("MLB RSN List", "Kumho list"):
            continue
        if sname not in SHEET_TO_HOME:
            log.append((sname, "skipping — not in SHEET_TO_HOME"))
            continue
        print(f"\n== {sname} ==")
        ws = wb[sname]
        try:
            c, s, _ = parse_sheet(ws, sname, existing, log)
        except Exception as e:
            log.append((sname, f"FATAL: {e}"))
            continue
        total_created += c
        total_skipped += s
        print(f"  -> created={c} skipped={s}")

    print(f"\n=== TOTALS: created={total_created} skipped={total_skipped} ===")
    if log:
        print("\nLog:")
        for sn, msg in log:
            print(f"  [{sn}] {msg}")


if __name__ == "__main__":
    main()
